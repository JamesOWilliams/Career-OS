"""Focused regression tests for DOCX presentation tuning in build_docx().

Locks the rendering-only refinements that brought USAA/HEB DOCX to ~2 pages:
tight margins, exact ~11.5pt leading, small paragraph spacing, and
keep-with-next chaining on section headings / roles / dates. Content and
reading order are covered elsewhere (backend_test.py)."""
import io
from docx import Document
from docx.shared import Pt, Inches
from resume_export import build_docx

REP = {
    "sections": {
        "identity": {
            "full_name": "Jane Doe",
            "email": "jane@example.com",
            "phone": "2105550251",
            "location": "Austin, TX",
            "links": [{"label": "GitHub", "url": "https://github.com/janedoe"}],
        },
        "summary": "Seasoned product leader with a track record of delivery.",
        "experiences": [
            {
                "track": "civilian", "role": "Product Manager", "employer": "Acme",
                "display_dates": "2020 – 2024",
                "bullets": [{"text": "Led a cross-functional team to ship a platform."},
                            {"text": "Grew adoption across multiple business units."}],
            },
        ],
    }
}


def _doc():
    return Document(io.BytesIO(build_docx(REP, "PM", "Acme")))


def test_margins_match_pdf_geometry():
    s = _doc().sections[0]
    assert s.top_margin == Inches(0.55)
    assert s.bottom_margin == Inches(0.5)
    assert s.left_margin == Inches(0.7)
    assert s.right_margin == Inches(0.7)


def test_normal_style_is_tight():
    styles = _doc().styles
    npf = styles["Normal"].paragraph_format
    assert npf.line_spacing == Pt(11.5)
    assert npf.space_after == Pt(2)
    assert styles["Normal"].font.size == Pt(10)
    lb = styles["List Bullet"].paragraph_format
    assert lb.line_spacing == Pt(11.5)
    assert lb.space_after == Pt(1)


def test_keep_with_next_chain_present():
    doc = _doc()
    # At least one heading and one meta/role paragraph must carry keep_with_next
    # so entries do not orphan their heading at a page bottom.
    kept = [p for p in doc.paragraphs if p.paragraph_format.keep_with_next]
    assert len(kept) >= 3, "expected heading/role/date keep-with-next chaining"
    # Name must not be clipped: it keeps multiple (non-exact) line spacing.
    name_p = doc.paragraphs[0]
    assert name_p.runs[0].font.size == Pt(18)
    assert name_p.paragraph_format.line_spacing != Pt(11.5)
